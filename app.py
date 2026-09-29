import os
import re
import streamlit as st
from groq import Groq
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document

load_dotenv()

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")

MAX_TOKENS = 800
MAX_HISTORY_MESSAGES = 12

ASSISTANT_AVATAR = "✅"
USER_AVATAR = "🧑‍💼"

GITHUB_URL = "https://github.com/OnkarDureja"
LINKEDIN_URL = "https://www.linkedin.com/in/onkar-dureja02/"
CODOLIO_URL = "https://codolio.com/profile/onkardureja"
RESUME_URL = "https://drive.google.com/file/d/1iGS2he3qsJ84WhdjySyDPPthXWrTJqUB/view?usp=drive_link"

st.set_page_config(page_title="AskCV — Onkar Dureja", page_icon="✅")

CUSTOM_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Source+Serif+4:wght@600;700&family=IBM+Plex+Sans:wght@400;500;600&display=swap');

html, body, [class*="css"], .stApp {
    font-family: 'IBM Plex Sans', sans-serif;
}

.askcv-header {
    text-align: center;
    padding: 1.25rem 0 0.5rem 0;
}

.askcv-title {
    font-family: 'Source Serif 4', serif;
    font-weight: 700;
    font-size: 2.6rem;
    color: inherit;
    margin-bottom: 0.2rem;
}

.askcv-tagline {
    font-size: 1rem;
    color: inherit;
    opacity: 0.7;
    margin-bottom: 0.85rem;
}

.askcv-badge {
    display: inline-block;
    background-color: rgba(16, 158, 110, 0.14);
    color: #109E6E;
    font-size: 0.82rem;
    padding: 0.35rem 0.9rem;
    border-radius: 999px;
    font-weight: 500;
}

.askcv-score-wrap {
    display: flex;
    align-items: baseline;
    gap: 0.6rem;
    margin: 1rem 0 0.5rem 0;
}

.askcv-score-label {
    font-size: 0.9rem;
    font-weight: 600;
    color: inherit;
    opacity: 0.7;
}

.askcv-score-value {
    font-family: 'Source Serif 4', serif;
    font-weight: 700;
    font-size: 2.4rem;
    line-height: 1;
}

.askcv-score-max {
    font-size: 1rem;
    color: inherit;
    opacity: 0.6;
}

.askcv-sidebar-link {
    display: block;
    padding: 0.45rem 0;
    color: inherit;
    text-decoration: none;
    font-size: 0.92rem;
    border-bottom: 1px solid rgba(128, 128, 128, 0.25);
}

.askcv-sidebar-link:hover {
    color: #109E6E;
}

[data-testid="stChatInput"] {
    border-radius: 14px;
}

div.stButton > button {
    border-radius: 10px;
    font-size: 0.85rem;
    padding: 0.6rem 0.8rem;
    text-align: left;
    width: 100%;
}
</style>
"""

st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

SCORE_PATTERN = re.compile(r"\*\*Fit Score:\*\*\s*(\d{1,3})\s*/\s*100")


def score_color(score: int) -> str:
    if score >= 85:
        return "#109E6E"
    if score >= 60:
        return "#D89B12"
    return "#E2603B"


def render_response(text: str):
    match = SCORE_PATTERN.search(text)
    if not match:
        st.markdown(text)
        return

    score = int(match.group(1))
    before = text[: match.start()]
    after = text[match.end():]

    if before.strip():
        st.markdown(before)

    st.markdown(
        f"""
        <div class="askcv-score-wrap">
            <span class="askcv-score-label">Fit Score</span>
            <span class="askcv-score-value" style="color:{score_color(score)}">{score}</span>
            <span class="askcv-score-max">/ 100</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if after.strip():
        st.markdown(after)


@st.cache_resource
def get_client():
    return Groq(api_key=GROQ_API_KEY)


@st.cache_data
def load_portfolio_data():
    with open("portfolio_data.md", "r", encoding="utf-8") as f:
        return f.read()


def build_system_prompt():
    portfolio_data = load_portfolio_data()
    return f"""You are AskCV, a conversational assistant that answers recruiter questions about Onkar Dureja based strictly on the portfolio data below.

Hard rule: only answer using information present in the portfolio data. If something is not covered there, say plainly that this isn't part of Onkar's current portfolio data. Never guess, estimate, or infer skills, experience, or numbers that are not explicitly stated.

Keep answers concise and professional. You are speaking to recruiters and hiring managers.

Job description matching:
If the user provides a job description (pasted as text, or extracted from an uploaded PDF/DOCX file) and asks for a match, evaluation, or fit assessment against Onkar's background, respond using exactly this structure:

**Matched Skills**
- Short, specific skills or experience from the portfolio data that are relevant to this role. Use short skill names, not sentences. Do not copy phrasing from the job description itself.

**Missing Skills**
- Specific required skills or experience the portfolio data shows no evidence of. Only list things that genuinely matter for this role. If nothing important is missing, say "No major gaps found."

**Fit Score:** X/100
Give a holistic score reflecting overall fit for this specific role, based on depth and relevance of experience, project quality, and how well the background lines up with what the role actually needs. Do not compute this as a simple ratio of matched to unmatched requirements. Reserve 85 and above for genuinely strong fits, and go below 40 for weak fits. Two different job descriptions should rarely produce the same score unless the fit is genuinely comparable.

**Verdict**
One or two sentences on overall fit and the strongest reason for or against this being a good match.

Ground every part of this evaluation strictly in the portfolio data above. Do not invent skills or experience not stated there.

--- PORTFOLIO DATA START ---
{portfolio_data}
--- PORTFOLIO DATA END ---
"""


def read_pdf_text(file_obj) -> str:
    reader = PdfReader(file_obj)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx_text(file_obj) -> str:
    doc = Document(file_obj)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + "\n"
    return text


def extract_text_from_upload(uploaded_file) -> str:
    suffix = uploaded_file.name.lower().rsplit(".", 1)[-1] if "." in uploaded_file.name else ""
    if suffix == "pdf":
        text = read_pdf_text(uploaded_file)
    elif suffix == "docx":
        text = read_docx_text(uploaded_file)
    else:
        raise ValueError(f"Unsupported file type: {uploaded_file.name}")

    if not text.strip():
        raise ValueError(f"No extractable text found in: {uploaded_file.name} (likely a scanned/image-only file)")
    return text


def handle_user_message(typed_text: str, uploaded_files=None):
    uploaded_files = uploaded_files or []
    extracted_chunks = []
    file_names = []
    extraction_errors = []

    for f in uploaded_files:
        try:
            extracted_chunks.append(extract_text_from_upload(f))
            file_names.append(f.name)
        except ValueError as e:
            extraction_errors.append(str(e))

    for err in extraction_errors:
        st.error(err)

    if not typed_text and not extracted_chunks:
        return

    if extracted_chunks:
        joined_extracted = "\n\n".join(extracted_chunks)
        content_for_model = (
            (typed_text + "\n\n" if typed_text else "")
            + f"[Content extracted from uploaded file(s): {', '.join(file_names)}]\n{joined_extracted}"
        )
        display_text = typed_text if typed_text else f"📎 Uploaded: {', '.join(file_names)}"
    else:
        content_for_model = typed_text
        display_text = typed_text

    st.session_state.messages.append({
        "role": "user",
        "content": content_for_model,
        "display": display_text,
    })
    with st.chat_message("user", avatar=USER_AVATAR):
        st.markdown(display_text)

    history = st.session_state.messages[-MAX_HISTORY_MESSAGES:]
    api_history = [{"role": m["role"], "content": m["content"]} for m in history]
    api_messages = [{"role": "system", "content": build_system_prompt()}] + api_history

    with st.chat_message("assistant", avatar=ASSISTANT_AVATAR):
        placeholder = st.empty()
        full_response = ""
        try:
            stream = client.chat.completions.create(
                model=GROQ_MODEL,
                messages=api_messages,
                max_tokens=MAX_TOKENS,
                stream=True,
            )
            for chunk in stream:
                delta = chunk.choices[0].delta.content
                if delta:
                    full_response += delta
                    placeholder.markdown(full_response + "▌")
            placeholder.empty()
            render_response(full_response)
        except Exception as e:
            full_response = f"Something went wrong talking to the model: {e}"
            placeholder.markdown(full_response)

    st.session_state.messages.append({"role": "assistant", "content": full_response})


client = get_client()

with st.sidebar:
    st.markdown("### Onkar Dureja")
    st.caption("B.Tech CSE (AI & ML), 2027 · iOS and AI engineering")
    st.markdown(
        f"""
        <a class="askcv-sidebar-link" href="{RESUME_URL}" target="_blank">Resume</a>
        <a class="askcv-sidebar-link" href="{GITHUB_URL}" target="_blank">GitHub</a>
        <a class="askcv-sidebar-link" href="{LINKEDIN_URL}" target="_blank">LinkedIn</a>
        <a class="askcv-sidebar-link" href="{CODOLIO_URL}" target="_blank">Codolio</a>
        """,
        unsafe_allow_html=True,
    )
    st.write("")
    if st.button("Clear chat", key="clear_chat"):
        st.session_state.messages = []
        st.rerun()

st.markdown(
    """
    <div class="askcv-header">
        <div class="askcv-title">AskCV</div>
        <div class="askcv-tagline">Ask about Onkar's background, or drop a job description for a fit check.</div>
        <span class="askcv-badge">✓ Answers grounded only in verified portfolio data</span>
    </div>
    """,
    unsafe_allow_html=True,
)

if "messages" not in st.session_state:
    st.session_state.messages = []

for message in st.session_state.messages:
    avatar = ASSISTANT_AVATAR if message["role"] == "assistant" else USER_AVATAR
    with st.chat_message(message["role"], avatar=avatar):
        if message["role"] == "assistant":
            render_response(message["content"])
        else:
            st.markdown(message.get("display", message["content"]))

pending_text = None

if not st.session_state.messages:
    example_questions = [
        "What's Onkar's AI/ML experience?",
        "Walk me through his iOS projects",
        "How is this bot prevented from making things up?",
    ]
    cols = st.columns(3)
    for col, ex in zip(cols, example_questions):
        with col:
            if st.button(ex, key=ex):
                pending_text = ex

user_input = st.chat_input(
    "Ask a question, or paste/attach a job description...",
    accept_file=True,
    file_type=["pdf", "docx"],
)

if pending_text:
    handle_user_message(pending_text)
elif user_input:
    typed_text = user_input.text.strip() if user_input.text else ""
    uploaded_files = user_input.files if user_input.files else []
    handle_user_message(typed_text, uploaded_files)